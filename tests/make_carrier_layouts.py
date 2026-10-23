"""Generate FICTIONAL returns in the layouts US carriers use, for Spoor's format tests.

Only the layouts are realistic (column names, preambles, how cells and times are written); every value is
invented. Phone numbers are in the 555-01xx fiction range, IMSIs use the 001-01 test network, coordinates are
arbitrary points, and the carriers are named only generically in the file contents.

Layouts covered (one folder each, as a case would be organised):
  mediation/    'CDR_Mediations_<number>_<id>.xlsx': agency preamble, header on row 12, separate Date and Time,
                1st/Last LTE Site ID + Sector ID, tower coordinates on voice rows only (texts have no cell);
                target named only in the file name. Tower list with eNodeB_ID, LTE_CI and a ListDate column.
  volte/        VoLTE CSV with local 'Record Open Date/Time (GMT -7)' and 'Record Open Dt/Tm(GMT)', 'Cell ID' N/A,
                eNB ID + Cell Face, Dir MO/MT/MF, MSISDN / Called # / CPN; list with hex ECGI 'MCCMNC-xxxxxxx'.
  report/       tab-separated report text with two tables (voice, data) under 'Voice Usage For:' / 'Data Usage For:'
                labels, CellLocation '[ECI:eNB:lon:lat:az:...]', one IMEI damaged into scientific notation.
  hda/          xlsx with an 'o Target MSISDN :' preamble, a date column at midnight plus an '(time)' column.
  dumps/        tower dumps at two scenes: a comma text file with one 'Search Criteria' section per scene
                (Number Associated w/Tower, ECGI 'first:last') and a VoLTE-layout dump per scene. One phone
                is at both scenes; one handset appears at both scenes under two numbers.
  location/     'TimingAdvance_<IMSI>_...xlsx' (End_Latitude/End_Longitude/Confidence, 15-digit Start_Cell)
                whose IMSI appears in the mediation CDR, and a pings sheet (Latitude/Longitude/TimeUtc/Conf/
                PhoneNumber). Copied next to the mediation CDR in the tests.

usage: python make_carrier_layouts.py OUT_DIR
"""
from __future__ import annotations

import csv
import datetime as dt
import random
import sys
from pathlib import Path

R = random.Random(20240521)
UTC = dt.timezone.utc

MED_T = "2025550111"          # mediation target (file name only)
MED_IMSI = "001010000000111"
MED_IMEI = "35391110000001"   # 14 digits + '0' as carriers often write it
VZ_T = "5715550122"
RPT_T = "(703)555-0133"
HDA_T = "12405550144"
BOTH = "+17035550155"         # in the dumps at both scenes
TWIN = ("+15715550166", "+15715550167")   # one handset, two numbers, both scenes
TWIN_IMEI = "99000000000077"
MCCMNC = "310260"

# (eNB, cell, lat, lon, az, bw)
SITES = [(48221, 2, 38.80010, -77.30020, 110, 65), (48221, 3, 38.80010, -77.30020, 230, 65),
         (48230, 7, 38.81020, -77.31030, 0, 65), (49102, 13, 38.82030, -77.32040, 120, 65)]
VZ_SITES = [(51234, 2, 38.90010, -77.10020, 120, 60), (51234, 22, 38.90010, -77.10020, 240, 60),
            (51250, 12, 38.91020, -77.11030, 0, 60)]


_USED: set = set()


def other(area: str = "") -> str:
    """A fictional 555-01xx number not used before in this run (so only planted overlaps exist)."""
    while True:
        n = f"1{area or R.choice(['703', '571', '202', '301', '240', '540', '804', '410'])}55501{R.randint(0, 99):02d}"
        if n not in _USED and n[1:] not in (BOTH[2:], TWIN[0][2:], TWIN[1][2:], MED_T, VZ_T, HDA_T[1:], "7035550133"):
            _USED.add(n)
            return n


def mediation(out: Path):
    import openpyxl
    d = out / "mediation"
    d.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1. Call Data"
    pre = [["Information Provided To: "], [], ["Agency:", "", "", "FICTIONAL DEMO"], ["Requestor:", "", "", "Training"],
           ["Agent Address:", "", "", "Training"], ["Billing City, State, Zip:", "", "", "Training"], ["Provided On:", "X"], [],
           ["This is in response to the legal process. All times below are reflected in Coordinated Universal Time (UTC)."], [], []]
    for r in pre:
        ws.append(r)
    ws.append(["Date", "Time", "Duration", "Call Type", "Direction", "Calling Number", "Dialed Number", "Called Number",
               "Destination Number", "IMSI", "IMEI", "Completion Code", "Service Code", "Switch Name", "1st LTE Site ID",
               "1st LTE Sector ID", "1st LAC", "1st Cell ID", "1st Tower Azimuth", "1st Tower LAT", "1st Tower LONG",
               "1st Tower Address", "Last LTE Site ID", "Last LTE Sector ID"])
    t0 = dt.datetime(2024, 5, 18, 0, 30, 0)
    for i in range(40):
        t = t0 + dt.timedelta(hours=3 * i, minutes=R.randint(0, 50))
        o = other()
        enb, cell, la, lo, az, _bw = SITES[i % len(SITES)]
        if i % 4 == 3:       # a text: no cell
            ws.append([t, t, "", "SMSC", "", o, "", "1" + MED_T, "", "", "", "Completed Successfully", "", "smsc01"] + [""] * 10)
            continue
        out_ = i % 2 == 0
        inline = i % 3 != 0  # some voice rows give the tower inline, others only the eNB + sector
        ws.append([t, t, R.randint(5, 600), "moc" if out_ else "mtc", "Outgoing" if out_ else "Incoming",
                   "1" + MED_T if out_ else o, o if out_ else "1" + MED_T, o if out_ else "1" + MED_T, o if out_ else "1" + MED_T,
                   MED_IMSI, MED_IMEI + "0", "Completed Successfully", "", "TAS201", enb, cell, "", "",
                   az if inline else "", la if inline else "", lo if inline else "", "1 Fictional Way" if inline else "",
                   enb, cell])
    wb.save(d / f"CDR_Mediations_{MED_T}_8592849.xlsx")
    with open(d / "Tower_List_LTE.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Region", "Market", "eNodeB_ID", "Sector_ID", "Latitude", "Longitude", "Azimuth", "BeamWidth", "TAC", "LTE_CI",
                    "Type", "ListDate"])
        for enb, cell, la, lo, az, bw in SITES:
            w.writerow(["EAST", "FICTION", enb, f"FAKE{enb}A{cell // 10 or 1}{cell % 10}", la, lo, az, bw, 10101, enb * 256 + cell,
                        "Vendor", "7/1/2022"])
        for k in range(20):   # unused sites
            w.writerow(["EAST", "FICTION", 60000 + k, f"FAKE{60000 + k}A11", 38.5 + k / 100, -77.5, 0, 65, 10101, (60000 + k) * 256 + 1,
                        "Vendor", "7/1/2022"])


def volte(out: Path):
    import openpyxl
    d = out / "volte"
    d.mkdir(parents=True, exist_ok=True)
    with open(d / f"CellSites-CDR_Historical_4g_VoLTE__{VZ_T}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Search MTN", "Search Start Date", "Search End Date", "Record Open Date/Time", "Record Open Dt/Tm(GMT)", "SID", "NID",
                    "Cell ID", "Cell Face", "Market ID", "eNB ID", "Dir", "MSISDN", "Called #", "CPN", "SOU", "RAT", "ENDPOINT", "VZWNE/V4B"])
        t0 = dt.datetime(2024, 5, 1, 16, 27, 50)
        for i in range(30):
            g = t0 + dt.timedelta(hours=5 * i)
            loc = g - dt.timedelta(hours=7)
            enb, cell, *_ = VZ_SITES[i % len(VZ_SITES)]
            o = other()
            dirn = ["MO", "MT", "MT"][i % 3]
            rows = [dirn] + (["MF"] if i % 6 == 2 else [])
            for dd in rows:
                called, cpn = (o, "1" + VZ_T) if dd == "MO" else ("1" + VZ_T, o)
                w.writerow([VZ_T, "5/1/24", "5/24/24", loc.strftime("%-m/%-d/%Y  %-I:%M:%S %p") + "(GMT -7)", g.strftime("%-m/%-d/%y %-H:%M"),
                            "N/A", "N/A", "N/A", cell, str(enb)[:2], enb, dd, VZ_T, called, cpn, R.randint(2, 900), "LTE", "N/A", "N/A"])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FICTION_CALEA.V0"
    ws.append(["Market SID", "Switch Number", "Switch Name", "Site Number", "E-911 Latitude Degrees (NAD83)",
               "E-911 Longitude Degrees (NAD83)", "Street Address", "City", "State", "Zip Code", "Sector", "Azimuth (deg)",
               "Antenna H-BW (deg)", "PN Offset", "Extended Base ID", "Channel Type", "eNodeB ID", "Vendor", "PCI", "ECGI"])
    for enb, cell, la, lo, az, bw in VZ_SITES:
        ws.append([None, None, None, 94, la, lo, "2 Fictional Ave", "Fiction", "VA", 20000, cell // 10 or cell, az, bw, None, None,
                   "LTE", enb, "Vendor", 289, f"311480-{enb * 256 + cell:07x}"])
    wb.save(d / "Fiction_Cell_site_list.xlsx")


def report(out: Path):
    d = out / "report"
    d.mkdir(parents=True, exist_ok=True)
    pad = "\t" * 6
    L = ["      MOBILITY WITH CELL LOCATION" + pad, "      " + pad,
         "      The carrier has queried for records from 05/24/2024 07:00:00am to 06/18/2024 11:59:00pm" + pad,
         "      Records are stored and provided in UTC." + pad, "      FICTIONAL DEMO RETURN" + pad, "      Matter ID:      0000001" + pad,
         "" + pad, "      Voice Usage For: " + RPT_T + pad,
         "      Item\tConnDateTime(UTC)\tSeizureTime\tET\tOriginatingNumber\tTerminatingNumber\tIMEI\tIMSI\tCT\tFeature\tCellLocation"]
    t0 = dt.datetime(2024, 5, 26, 12, 23)
    for i in range(8):
        t = t0 + dt.timedelta(hours=7 * i)
        o = other()
        out_ = i % 2 == 0
        imei = "3.50000E+14" if i == 5 else "35000000000133" + "0"
        loc = ("[000000020:0:0.0:0.0:-1:-1.0:]" if i % 4 == 3 else
               f"[{48221 * 256 + 2}:48221:-77.30020:38.80010:110:-1.0:200;{48230 * 256 + 7}:48230:-77.31030:38.81020:0:-1.0:200]")
        L.append(f"{i + 1}\t{t.month}/{t.day}/{t.year} {t.hour}:{t.minute:02d}\t0:07\t{R.randint(0, 9)}:{R.randint(10, 59)}\t"
                 f"{'17035550133' if out_ else o}\t{o if out_ else '17035550133'}\t{imei}\t001010000000133\t{'MO' if out_ else 'MT'}\t\t{loc}")
    L += ["      Data Usage For: " + RPT_T + pad,
          "      Item\tConnDateTime(UTC)\tOriginatingNumber\tET\tBytesUp\tBytesDn\tApn\tIMEI\tIMSI\tCellLocation" + "\t"]
    for i in range(4):
        t = t0 + dt.timedelta(hours=5 * i, minutes=30)
        L.append(f"{i + 1}\t{t.month}/{t.day}/{t.year} {t.hour}:{t.minute:02d}\t17035550133\t10:00\t2490\t2829\tims\t350000000001330\t"
                 f"001010000000133\t[{49102 * 256 + 13}:49102:-77.32040:38.82030:120:-1.0:300]")
    L += ["      FICTIONAL - authorized persons only" + pad]
    (d / "ReportAU_0000001.txt").write_text("\n".join(L) + "\n", encoding="utf-8")


def hda(out: Path):
    import openpyxl
    d = out / "hda"
    d.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Volte"
    for r in [["*" * 40], ["o Query type : volte"], [f"o Target MSISDN : {HDA_T}"], ["o Time Zone  : EST"],
              ["o Date range : [ 05-24-2024 00:00:00 to 05-24-2024 21:00:00 ]"], ["*" * 40]]:
        ws.append(r)
    ws.append(["start_report_date", "est (time)", "imsi", "device_identifier", "description", "direction", "calling_number",
               "called_number", "completion_code", "duration", "site_number", "enodeb_id", "sector_azimuth", "tower_lat", "tower_long",
               "tower_address", "tower_city", "tower_state", "tower_zip"])
    for i in range(20):
        tt = dt.time(1 + i, R.randint(0, 59), R.randint(0, 59))
        o = int(other())
        out_ = i % 2 == 0
        ws.append([dt.datetime(2024, 5, 24), tt, 1010000000144, 3500000000014410, "AUDIO CALL",
                   "Outgoing VoLTE Call" if out_ else "Incoming VoLTE Call", int(HDA_T) if out_ else o, o if out_ else int(HDA_T),
                   "SUCCESS", R.randint(5, 300), None, 123401, 250, 43.20001, -70.80002, "1 Fiction Rd", "Fiction", "ME", 3999])
    wb.save(d / f"0000001_CXM_HDA_volte_{HDA_T}_05242024000000_EST-Final.xlsx")


def dumps(out: Path):
    import openpyxl
    d = out / "dumps"
    d.mkdir(parents=True, exist_ok=True)
    scenes = [dt.datetime(2024, 6, 2, 7, 0, 0), dt.datetime(2024, 6, 5, 3, 0, 0)]
    L = []
    for k, s0 in enumerate(scenes):
        L += [f"Search Criteria - USID:00000{k + 1}",
              f"Date Range - From: {s0:%m/%d/%Y} 2:00 AM To: {s0:%m/%d/%Y} 3:00 AM Central",
              "Source,Conn Date,Conn Time (UTC),Disc Date,Disc Time (UTC),Originating Number,Dialed Number,Terminating Number,"
              "Call Type,Number Associated w/Tower,IMSI Associated w/Tower,IMEI Associated w/Tower,Minutes,Tower LAC/CID or ECGI,APN,MMS Type"]
        rows = [(BOTH[2:], "001010000000155", "35000000000155" + "10")] + [(TWIN[k][2:], f"00101000000016{6 + k}", TWIN_IMEI + "10")] \
            + [(other(), f"0010100000{R.randint(10000, 99999)}", f"35{R.randint(10 ** 11, 10 ** 12 - 1)}10") for _ in range(12)]
        for j, (num, imsi, imei) in enumerate(rows):
            t = s0 + dt.timedelta(minutes=2 * j + 1)
            src = ["Voice", "Data", "SMS"][j % 3]
            o = other()
            mo = j % 2 == 0
            L.append(f"{src},{t:%Y-%m-%d},{t:%H:%M:%S},{t:%Y-%m-%d},{t:%H:%M:%S},{num if mo else o},{o if mo else ''},{o if mo else num},"
                     f"{'MO' if mo else 'MT'},{num},{imsi},{imei},0.5,\"{12355555 + k}:{12355555 + k}\",{'ims' if src == 'Data' else ''},")
        L += ["", "Records show transactions processed through the dumped cells; FICTIONAL DEMO.", ""]
    (d / "ReportTOWER_0000001.txt").write_text("\n".join(L) + "\n", encoding="utf-8")
    for k, s0 in enumerate(scenes):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"flatlry_00000{k}"
        ws.append(["Record Open Date/Time", "Record Open Dt/Tm(GMT)", "SID", "NID", "Cell ID", "Cell Face", "Market ID", "eNB ID", "Dir",
                   "MSISDN", "Called #", "CPN", "SOU", "RAT", "ENDPOINT", "VZWNE/V4B", "User Equip Type"])
        for j in range(15):
            t = s0 + dt.timedelta(minutes=3 * j + 2)
            num = other()[1:] if j else BOTH[2:]
            ws.append([(t - dt.timedelta(hours=5)).strftime("%Y/%m/%d %H:%M:%S") + "(GMT  -5)", t.strftime("%Y-%m-%d %H:%M:%S"), "N/A", "N/A",
                       "N/A", "22", "612", f"61234{k}", "MO", num, "1" + other()[1:], "1" + num, R.randint(3, 400), "LTE", "N/A", "N/A", "null"])
        wb.save(d / f"VoLTE_CDRs_Scene {k + 1}, Fiction City.xlsx")


def location(out: Path):
    import openpyxl
    d = out / "location"
    d.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TimingAdvance"
    ws.append(["All times below are reflected in Coordinated Universal Time (UTC). FICTIONAL DEMO."])
    for _ in range(4):
        ws.append([])
    ws.append(["End_Latitude", "End_Longitude", "Confidence", "IMSI", "IMEI", "Model", "Make", "Service_Type", "Start_DateTime", "Duration",
               "End_DateTime", "Start_Cell", "Start_eNodeB", "Start_TAC_LAC", "Start_Sector", "End_Cell", "End_eNodeB", "End_Sector",
               "Start_Timing_Advance_Miles", "End_Timing_Advance_Miles", "Device_Technology", "Market"])
    t0 = dt.datetime(2024, 5, 20, 0, 0, 4)
    for i in range(30):
        enb, cell, la, lo, *_ = SITES[i % 2]
        t = t0 + dt.timedelta(minutes=17 * i)
        ws.append([f"{la + 0.003:.6f}", f"{lo + 0.002:.6f}", "Low", MED_IMSI, MED_IMEI + "3", "PHONE", "MAKER", "Data Only", t, "5",
                   t + dt.timedelta(seconds=5), f"{MCCMNC}{enb * 256 + cell:09d}", str(enb), "10101", cell,
                   f"{MCCMNC}{enb * 256 + cell:09d}", str(enb), cell, 0.29, 0.31, "LTE", "FICTION"])
    wb.save(d / f"TimingAdvance_{MED_IMSI}_05-20-2024_05-22-2024_Records.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(["Latitude", "Longitude", "LocalTime", "TimeUtc", "Conf", "PhoneNumber", "VehiclePlate", "EndDate"])
    for i in range(12):
        t = dt.datetime(2024, 5, 21, 5, 6, 21) + dt.timedelta(minutes=10 * i)
        ws.append([f"{38.8001 + i / 1000:.6f}", "-77.300200", (t - dt.timedelta(hours=5)).strftime("%b %d %Y %H:%M:%S"),
                   t.strftime("%b %d %Y %H:%M:%S"), "4043", f"{MED_T[:3]}-{MED_T[3:6]}-{MED_T[6:]}", "", ""])
    wb.save(d / "Fiction Pings.xlsx")


def main(out_dir: str):
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    mediation(out)
    volte(out)
    report(out)
    hda(out)
    dumps(out)
    location(out)
    (out / "README.txt").write_text(
        "FICTIONAL DEMO DATA generated by tests/make_carrier_layouts.py - carrier LAYOUTS with invented values only.\n",
        encoding="utf-8")


EXPECT = {"mediation_target": "+1" + MED_T, "volte_target": "+1" + VZ_T, "report_target": "+17035550133",
          "hda_target": "+" + HDA_T, "both_scenes": BOTH, "twin_numbers": set(TWIN), "twin_imei": TWIN_IMEI,
          "ta_imsi": MED_IMSI}

if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(Path(__file__).resolve().parent.parent / "demo" / "carrier_layouts"))
