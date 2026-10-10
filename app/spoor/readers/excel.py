"""XLSX/XLSM (openpyxl, read-only streaming) and legacy XLS (xlrd). Citations: Sheet!R<row>."""
from __future__ import annotations

import datetime as dt
from typing import Iterator, Optional

from ..ipaddr import split_ip_port
from ..timeparse import looks_like_timestamp
from .textmode import Rec


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    return str(v)


def _headerish(row) -> bool:
    vals = [v for v in row if v not in (None, "")]
    if len(vals) < 2:
        return False
    bad = 0
    for v in vals:
        if isinstance(v, (int, float, dt.datetime, dt.date)) and not isinstance(v, bool):
            bad += 1
        elif split_ip_port(str(v)) or looks_like_timestamp(v):
            bad += 1
    return bad <= len(vals) * 0.2


def _pick_header(rows: list[tuple]) -> tuple[Optional[int], int]:
    """(header_index or None, ncols) from the first rows of a sheet."""
    widths = [sum(1 for v in r if v not in (None, "")) for r in rows]
    if not widths or max(widths) < 2:
        return None, max(widths or [1])
    target = max(widths)
    for i, r in enumerate(rows[:40]):
        if widths[i] >= max(2, target * 0.6) and _headerish(r):
            later = widths[i + 1: i + 21]
            if later and sum(1 for w in later if w >= 2) >= max(1, len(later) * 0.5):
                return i, len(r)
    return None, max(len(r) for r in rows)


def _headers_from(row, ncols) -> list[str]:
    out, seen = [], {}
    for i in range(ncols):
        v = row[i] if i < len(row) else None
        h = _cell_text(v).strip() or f"col{i + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h}_{seen[h]}"
        else:
            seen[h] = 1
        out.append(h)
    return out


# ---------------------------------------------------------------- xlsx

def _open_xlsx(path):
    import openpyxl
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def scan_xlsx(path, name: str) -> dict:
    wb = _open_xlsx(path)
    out = {"kind": "xlsx", "encoding": None, "tables": [], "preamble": [], "notes": []}
    try:
        for ws in wb.worksheets:
            head = []
            for r in ws.iter_rows(values_only=True, max_row=260):
                head.append(tuple(r))
            if not any(any(v not in (None, "") for v in r) for r in head):
                out["notes"].append(f"sheet '{ws.title}' is empty")
                continue
            hidx, ncols = _pick_header(head)
            state = getattr(ws, "sheet_state", "visible")
            if state != "visible":
                out["notes"].append(f"sheet '{ws.title}' is {state} in the workbook - included")
            if hidx is None:
                headers = [f"col{i + 1}" for i in range(ncols)]
                start = 0
            else:
                headers = _headers_from(head[hidx], ncols)
                start = hidx + 1
            pre = [(f"{ws.title}!R{i + 1}", " | ".join(_cell_text(v) for v in r if v not in (None, "")))
                   for i, r in enumerate(head[:start - 1 if hidx is not None else 0])]
            out["preamble"] += [p for p in pre if p[1]]
            samples = [[_cell_text(v) for v in r[:ncols]] for r in head[start:start + 200]
                       if any(v not in (None, "") for v in r)]
            raw_samples = [list(r[:ncols]) for r in head[start:start + 200] if any(v not in (None, "") for v in r)]
            try:
                dims = ws.max_row
            except Exception:
                dims = None
            out["tables"].append({
                "name": ws.title, "mode": "table", "header_line": (hidx + 1) if hidx is not None else None,
                "data_start_line": start + 1, "headers": headers, "ncols": ncols, "samples": samples,
                "raw_samples": raw_samples, "rows_estimate": dims,
                "header_cite": f"{ws.title}!R{hidx + 1}" if hidx is not None else None,
            })
            out["notes"].append(f"sheet '{ws.title}': {ncols} columns, header "
                                + (f"at row {hidx + 1}" if hidx is not None else "not found") +
                                (f", about {dims:,} rows" if dims else ""))
    finally:
        wb.close()
    return out


def iterate_xlsx(path, cfg: dict, table: dict, resume: Optional[dict] = None) -> Iterator[tuple[Rec, dict]]:
    wb = _open_xlsx(path)
    try:
        ws = wb[table["name"]]
        headers, ncols, start = table["headers"], table["ncols"], table["data_start_line"]
        skip = (resume or {}).get("rows", 0)
        rows = 0
        for ridx, r in enumerate(ws.iter_rows(values_only=True, min_row=start), start=start):
            vals = list(r[:ncols]) + [None] * max(0, ncols - len(r))
            if not any(v not in (None, "") for v in vals):
                continue
            rows += 1
            if rows <= skip:
                continue
            raw = " | ".join(_cell_text(v) for v in vals)
            yield Rec(f"{ws.title}!R{ridx}", ridx, None, None, raw, dict(zip(headers, vals)), ws.title), {"rows": rows}
    finally:
        wb.close()


# ---------------------------------------------------------------- xls

def _xls_rows(sh, book):
    import xlrd
    for i in range(sh.nrows):
        row = []
        for c in sh.row(i):
            if c.ctype == xlrd.XL_CELL_DATE:
                try:
                    row.append(xlrd.xldate.xldate_as_datetime(c.value, book.datemode))
                except Exception:
                    row.append(c.value)
            elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                row.append(None)
            else:
                row.append(c.value)
        yield i + 1, tuple(row)


def scan_xls(path, name: str) -> dict:
    import xlrd
    book = xlrd.open_workbook(path, on_demand=True)
    out = {"kind": "xls", "encoding": None, "tables": [], "preamble": [], "notes": []}
    try:
        for sh in book.sheets():
            head = [row for _, (_, row) in zip(range(260), _xls_rows(sh, book))]
            if not head:
                continue
            hidx, ncols = _pick_header(head)
            headers = _headers_from(head[hidx], ncols) if hidx is not None else [f"col{i + 1}" for i in range(ncols)]
            start = (hidx + 1) if hidx is not None else 0
            out["preamble"] += [(f"{sh.name}!R{i + 1}", " | ".join(_cell_text(v) for v in r if v not in (None, "")))
                                for i, r in enumerate(head[:max(0, start - 1)])]
            out["tables"].append({
                "name": sh.name, "mode": "table", "header_line": (hidx + 1) if hidx is not None else None,
                "data_start_line": start + 1, "headers": headers, "ncols": ncols,
                "samples": [[_cell_text(v) for v in r[:ncols]] for r in head[start:start + 200]],
                "raw_samples": [list(r[:ncols]) for r in head[start:start + 200]], "rows_estimate": sh.nrows,
                "header_cite": f"{sh.name}!R{hidx + 1}" if hidx is not None else None,
            })
    finally:
        book.release_resources()
    return out


def iterate_xls(path, cfg: dict, table: dict, resume: Optional[dict] = None) -> Iterator[tuple[Rec, dict]]:
    import xlrd
    book = xlrd.open_workbook(path, on_demand=True)
    try:
        sh = book.sheet_by_name(table["name"])
        headers, ncols, start = table["headers"], table["ncols"], table["data_start_line"]
        skip = (resume or {}).get("rows", 0)
        rows = 0
        for ridx, r in _xls_rows(sh, book):
            if ridx < start:
                continue
            vals = list(r[:ncols]) + [None] * max(0, ncols - len(r))
            if not any(v not in (None, "") for v in vals):
                continue
            rows += 1
            if rows <= skip:
                continue
            raw = " | ".join(_cell_text(v) for v in vals)
            yield Rec(f"{sh.name}!R{ridx}", ridx, None, None, raw, dict(zip(headers, vals)), sh.name), {"rows": rows}
    finally:
        book.release_resources()
